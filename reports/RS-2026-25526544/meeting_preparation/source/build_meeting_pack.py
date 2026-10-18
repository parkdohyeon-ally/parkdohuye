from pathlib import Path
import asyncio, html, json, zipfile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from playwright.async_api import async_playwright
import fitz

ROOT = Path(__file__).resolve().parents[1]
E = html.escape
ROWS = [
 ('D1', '이용 대상·기능', '회원등급별 DT 접속, 직책·임무별 웹 열람 권한 (28–29, 34, 63쪽)', '내부 연구자·외부 참여자·회원은 어느 화면에서 무엇을 이용하는가?', '사용자 × 검색·DT·예측·제어 범위', '조합·인실리코·이즈파크·세부기관'),
 ('D2', '데이터 제공 범위', '일반·보안 데이터 분리, 추가 인증을 적용한 실데이터 제공 (23, 29, 63쪽)', '제공할 데이터·원본/사본·출처·갱신주기는? 공개 범위는 누가 확인하는가?', '데이터 목록·출처·제공 대상·확인 주체', '세부기관·인실리코·이즈파크·조합'),
 ('D3', '망 격리·연결', '총괄–세부 송수신, 충분한 데이터 확보 후 로컬 모드 (63쪽)', '외부→내부 접속 또는 내부→외부 전송이 가능한가? VPN·중계·사본 중 가능한 방식은?', '송신 시작점·수신 시스템·허용 경로', 'KETI·세부기관·각 시스템/망 담당'),
 ('D4', '인증·접근권한', '조합의 인증·접근통제, 인실리코의 권한별 웹·API (28–29, 34쪽)', '회원·등급·추가 인증과 최종 데이터 권한 확인은 각각 어디서 수행하는가?', '담당 시스템·전달 정보·회수·기록 방식', '조합·인실리코·이즈파크·세부기관'),
 ('D5', '기능·연계 담당', '미들웨어 수집, DT 송수신·저장, DT↔예측 API (26, 28, 35, 46쪽)', '수집·저장·검색·예측·시각화 연결별 개발 담당과 입력·출력은?', '기능별 담당·API/데이터 항목·시험 담당', 'KETI·한양대·성균관대·인실리코·이즈파크·조합'),
 ('D6', '실행·비용·검증', '총괄 모듈별 역할 정의, 조합 핵심 R&D 자체 개발·외주 범위 (21, 28–29쪽)', '재사용 자원과 추가 개발은? 누가 언제까지 무엇으로 검증하는가?', '담당기관·기한·비용 확인·완료 증빙', 'KETI·모든 관련기관'),
]

CSS = '''*{box-sizing:border-box}body{margin:0;color:#183249;font-family:"Noto Sans CJK KR",sans-serif;background:white}h1{font-size:24px;margin:0 0 6px}h2{font-size:17px;margin:14px 0 7px;color:#08787e}p{margin:5px 0;line-height:1.5}table{width:100%;border-collapse:collapse;table-layout:fixed}th{background:#183249;color:white;text-align:left;padding:9px}td{border:1px solid #cbd7df;padding:8px;vertical-align:top;line-height:1.45}tbody tr:nth-child(even){background:#f2f7fa}.sub{color:#52697c}.tag{display:inline-block;color:#06777b;background:#e7f5f3;padding:4px 9px;font-weight:700;font-size:12px}.note{background:#fff6e4;border-left:4px solid #c58a20;padding:9px 12px;margin:10px 0}.band{background:#eaf4fa;padding:9px 12px;margin:10px 0}.small{font-size:11px}.muted{color:#52697c}.sheet{padding:28px 34px;width:1123px;min-height:794px}.sheet table{font-size:12px}.sheet td{height:76px}.sheet footer{font-size:10px;margin-top:9px;color:#52697c}.page{width:794px;height:1122px;padding:37px 43px;page-break-after:always;font-size:12px}.page:last-child{page-break-after:auto}.page h1{font-size:23px}.page table{font-size:11px}.page th,.page td{padding:7px}.page ul,.page ol{padding-left:20px;margin:5px 0}.page li{margin:4px 0;line-height:1.45}.quote{border-left:4px solid #08787e;padding:10px 12px;background:#edf7f6}.page footer{font-size:10px;color:#52697c;margin-top:14px}'''

def doc(body, title):
    return '<!doctype html><html lang="ko"><meta charset="utf-8"><title>'+E(title)+'</title><style>'+CSS+'.sheet{padding:22px 30px}.sheet td{height:66px}.sheet td,.sheet th{padding:7px}.page{height:1100px;break-inside:avoid}</style><body>'+body+'</body></html>'

def build_sheet():
    body='''<main class="sheet"><span class="tag">수요일 회의 · 협의안 / 결정 전</span><h1>AI 자율실험실 기관별 역할·연계 협의표</h1><p class="sub">목표: 데이터 보호와 권한별 이용 기능을 함께 구현할 범위·담당·검증 기준 확인</p><p class="small">일시: __________　 방식: □ 온라인 □ 오프라인　 참석: __________________　 기록: __________</p><div class="band small"><b>진행 원칙</b>　계획서 기재 내용 → 기관별 구현 조건 → 공동 결정 순으로 확인합니다. 역할 중복·서버 소유·통합 인증을 미리 확정하지 않습니다.</div><table><colgroup><col style="width:11%"><col style="width:23%"><col style="width:31%"><col style="width:18%"><col style="width:17%"></colgroup><thead><tr><th>협의 항목</th><th>계획서 기재 내용<br><span style="font-weight:400">PDF 파일 쪽수</span></th><th>회의에서 답할 질문</th><th>남길 결정사항</th><th>합의·미결 기록</th></tr></thead><tbody>'''
    for id,title,source,q,result,org in ROWS:
        body+=f'<tr><td><b>{id}<br>{E(title)}</b></td><td>{E(source)}</td><td>{E(q)}</td><td>{E(result)}</td><td class="muted">□합의 □추가확인<br>담당: __________<br>기한: __________</td></tr>'
    body+='''</tbody></table><div class="note small"><b>구분할 사항</b>　일반데이터 ≠ 누구나 공개하는 데이터 / 로컬 모드 ≠ 물리적 완전 망 분리 / VPN 접속 ≠ 데이터 열람 허용</div><p class="small"><b>종료 전 확인</b>　□ 합의 문구 공동 확인　□ 미결사항 담당·기한　□ 시범 연계 범위·증빙　□ 후속 확인 일정</p><footer>RS-2026-25526544 연구개발계획서 기반. 질문·검증·진행 방식은 회의 제안이며 계획서의 확정 설계가 아닙니다.<br>상세 결정문·참여기관 의견·담당·기한은 meeting_workbook.xlsx에 기록합니다. 기관명은 개발 담당을 뜻하며 서버 소유·운영 책임을 뜻하지 않습니다.</footer></main>'''
    (ROOT/'discussion_sheet.html').write_text(doc(body,'기관별 역할·연계 협의표'),encoding='utf-8')

PREP_MD = '''# 수요일 회의 준비·진행안

회의 일시·장소·방식·참석자는 미정입니다. 아래는 60분 회의 제안이며, 업무 배정이나 기관 간 합의가 완료되었다는 뜻이 아닙니다.

## 지금부터 준비할 일

| 시점 | 박도현 매니저 준비·확인 사항 | 준비 결과 |
|---|---|---|
| 지금 | 개최 시간·방식·참석 담당자 확인. 총괄·조합·인실리코·이즈파크를 중심으로 세부기관의 데이터/망 담당 참석 필요 여부 확인 | 참석·연락 명단 |
| 회의 전 | 각 기관에 망 격리 조건, 이용 대상, 제공 가능 데이터, 기존 서버·VPN·API 현황의 회신 요청 | 기관별 조건표 |
| 회의 전날까지(제안) | 회신 내용을 계획서 사실·기관 제안·미확인으로 구분하고 D1~D6에 연결 | 쟁점별 비교 메모 |
| 회의 전날까지(제안) | 협의표·근거 쪽수·대안 비교자료 공유, 예상 질문과 결정할 항목 안내 | 공통 검토자료 |
| 회의 직전 | 원문 PDF·협의표·기록용 엑셀 준비. 화면 공유와 기록 담당 확인 | 회의 준비 완료 |
| 회의 직후 | 합의 문구·미결 담당·기한을 정리해 공동 확인 요청 | 회의록·후속 조치표 |

## 기관별 사전 확인 질문

- KETI: 계획서 기능별 확인 담당과 기관 간 역할·연계 합의 절차는?
- 인실리코: 권한 적용 사용자는 누구이며, 외부 접속·외부 전송·기존 API 중 가능한 범위는?
- 이즈파크: DT 화면·저장소·예측 연동에 필요한 데이터와 접속 조건은?
- 세부기관: 공유·조회 허용 데이터, 공개 범위 확인 주체, 망 정책 결정 담당은?
- 조합: 대상 회원·등급·추가 인증·접속 대상 서비스·자체 개발 범위는?
- 한양대·성균관대: 장비 수집·표준 데이터·실험 운영 기능의 연계 지점과 시험 참여 범위는?

기관별 참석·역할은 제안입니다. 회신이 없으면 '미확인'으로 기록하며 반대 또는 동의로 해석하지 않습니다.

## 60분 진행안

| 시간 | 논의 | 기록할 결과 |
|---|---|---|
| 0~5분 | 목적과 사실·제안 구분 | 논의 범위 |
| 5~15분 | D1 이용 대상 / D2 제공 데이터 | 사용자·기능·데이터 범위 |
| 15~25분 | D3 격리 의미·통신 조건 | 가능한 송수신 경로 |
| 25~40분 | D4 인증 / D5 기능·연계 | 인증·권한·개발 경계 |
| 40~50분 | D6 비용·시험·일정 | 담당·기한·성공 기준 |
| 50~60분 | 합의 문구 재확인 | 결정·미결 목록, 다음 일정 |

오프닝: “데이터 보호와 계획서의 권한별 이용 기능을 함께 구현하기 위해, 각 기관의 전제를 확인하고 연결 지점과 담당 업무를 정리하고자 합니다.”

정리 질문: “오늘 정한 사항은 누구에게, 어떤 데이터·기능에 적용됩니까? 미정인 부분은 어느 기관이 언제까지 확인할 수 있을까요?”

## 선택지 비교 — 기술 대안, 견적·확정 설계 아님

| 대안 | 비용/난이도(상대) | 적용 조건·한계 |
|---|---|---|
| 공개자료 별도 게시 | 초기 낮음 / 난이도 낮음 | 이 기능만으로 권한별 DT·예측 연동 등을 대체하기 어려움 |
| VPN 지정 접속 | 초기 낮음~중간 / 난이도 중간 | 기존 VPN·소수 이용자 조건 확인. 서비스별 권한 별도 필요 |
| 제한 API 중계 | 초기·난이도 중간~높음 | 내부 API 접속 허용, 인증 연계와 보안 운영 담당 필요 |
| 허용 데이터 사본 제공 | 초기·난이도 중간~높음 | 외부 전송 허용 필요. 갱신·삭제·권한 반영, 예측·제어 별도 검토 |

세부 견적은 이용자 수·기관 수·데이터량·주기·기존 자원 확인 후 요청합니다. 조합의 핵심 연동·보안 R&D 자체 개발과 외주 범위(PDF 28~29쪽)를 함께 확인합니다.

## 회의 중 유용한 확인 문장

- 격리: “외부에서 들어오는 접속과 내부에서 내보내는 전송 중 어느 쪽이 제한됩니까?”
- 권한: “회원 인증과 최종 데이터 열람 판단이 같은 단계인지 구분해 보겠습니다.”
- 제공: “실데이터는 어느 저장소에서, 어떤 서비스로, 누구에게 제공됩니까?”
- 미결: “지금 결정하기 어려운 조건은 담당기관과 확인 기한을 남기겠습니다.”
- 역할: “개발 담당과 서버 운영 담당, 데이터 공개 범위 확인 담당을 각각 기록하겠습니다.”

## 시범 연계 검증안 — 제안, 아직 수행하지 않음

비민감 샘플 데이터로 다음 항목의 담당·기한·증빙을 먼저 합의합니다.

1. 허용 사용자: 필요한 인증 후 허용된 데이터 조회 성공.
2. 미인증/권한 부족 사용자: 보호된 데이터 요청 거부.
3. 권한 회수 사용자: 합의한 반영 시간 안에 접근 차단.
4. 기록: 사용자·시각·대상·허용/거부 결과를 정한 위치에서 확인.
5. 예측 연계: 포함하기로 한 경우 입력 조건과 결과의 대응 확인.
6. 조회용 사본: 채택한 경우 갱신·삭제·권한 변경의 반영 확인.

장비 제어는 조회와 구분해 대상·권한·시험 담당을 결정합니다. 조회 시범 검증만으로 전체 계획서 기능의 완료를 판단하지 않습니다.

## 회의가 끝날 때 남아야 할 것

- 합의한 범위와 적용 조건이 적힌 결정문
- 개발·운영·데이터 공개 범위 확인 담당을 구분한 역할표
- 송신 시작 시스템→수신 시스템→데이터/API→인증·권한 확인 경로
- 미결사항 담당·확인 기한, 시범 검증 담당·완료 증빙

## 계획서 근거

PDF 파일 페이지 기준. 문서에 인쇄된 쪽수와 구분합니다.

| PDF | 확인 내용 |
|---|---|
| 21 | 총괄 통합 아키텍처, 데이터 유입·유출, 모듈별 역할 정의 |
| 23 | 미들웨어→DB API, 일반·보안 데이터 분류 |
| 26, 28 | 표준 데이터 구조·파이프라인, 장비 데이터 수집·전달 |
| 28~29 | 조합 연동 규격·접근통제·2차 인증, 영상·실데이터, 자체 개발·외주 범위 |
| 34~35 | 사용자별 웹·API 권한, DT 실험 조건 입력·예측 결과 제공 |
| 46 | 이즈파크 수집 데이터 저장·대시보드·3D 시각화 |
| 63 | 총괄–세부 송수신, 보안데이터 분리, 충분한 데이터 확보 후 로컬 모드, 등급별 권한 |

기관별 이메일 의견은 계획서의 확정 내용과 구분하여 별도 확인합니다. 이 자료에는 메일 원문·연락처·미확정 기관 입장을 포함하지 않았습니다.
'''

def build_guide():
    (ROOT/'meeting_prep.md').write_text(PREP_MD,encoding='utf-8')
    p1='''<section class="page"><span class="tag">수요일 회의 준비안 · 60분 기준</span><h1>회의 전 준비와 당일 진행</h1><p>목표는 <b>적용 범위 → 연결 방식 → 담당·검증 기준</b>의 공동 확인입니다.<br>일시·방식·참석기관은 미정이며, 아래 일정과 배정은 제안입니다.</p><h2>지금부터 할 일</h2><table><colgroup><col style="width:19%"><col style="width:59%"><col style="width:22%"></colgroup><tr><th>시점</th><th>준비·확인</th><th>결과물</th></tr><tr><td>지금</td><td>개최 시간·방식·참석 담당자 확인. 세부기관의 데이터·망 담당 참여 필요 여부 확인</td><td>참석 명단</td></tr><tr><td>회의 전</td><td>기관별 망 조건·이용 대상·허용 데이터·기존 서버/VPN/API 현황 확인 요청</td><td>기관별 조건</td></tr><tr><td>전날까지<br>(제안)</td><td>회신을 계획서 사실·기관 제안·미확인으로 구분. 협의표와 근거 쪽수 공유</td><td>D1~D6 메모</td></tr><tr><td>회의 직전</td><td>원문 PDF·협의표·엑셀 준비. 화면 공유·기록 담당 확인</td><td>진행 준비</td></tr><tr><td>회의 직후</td><td>합의 문구·미결 담당·기한 공동 확인 요청</td><td>회의록</td></tr></table><h2>기관별로 미리 확인할 질문</h2><ul><li><b>KETI:</b> 계획서 기능·역할·연계 합의의 확인 절차는?</li><li><b>인실리코:</b> 권한 적용 대상, 외부 접속·전송, 기존 API 가능 범위는?</li><li><b>이즈파크:</b> DT 화면·저장소·예측 연동의 데이터·접속 조건은?</li><li><b>세부기관:</b> 제공 가능 데이터, 공개 범위 확인·망 정책 담당은?</li><li><b>조합:</b> 대상 회원·등급·추가 인증·접속 서비스·자체 개발 범위는?</li><li><b>한양대·성균관대:</b> 수집·표준 데이터·실험 운영 연계와 시험 범위는?</li></ul><h2>60분 회의 진행안</h2><table><tr><th style="width:18%">시간</th><th>논의</th><th style="width:32%">남길 결과</th></tr><tr><td>0~5분</td><td>목적과 사실·제안 구분</td><td>논의 범위</td></tr><tr><td>5~15분</td><td>D1 이용자 / D2 제공 데이터</td><td>사용자·기능·데이터</td></tr><tr><td>15~25분</td><td>D3 망 격리·통신 조건</td><td>가능한 송수신 경로</td></tr><tr><td>25~40분</td><td>D4 인증 / D5 기능·연계</td><td>권한·개발 경계</td></tr><tr><td>40~50분</td><td>D6 비용·시험·일정</td><td>담당·기한·성공 기준</td></tr><tr><td>50~60분</td><td>합의 문구 공동 확인</td><td>결정·미결·다음 일정</td></tr></table><h2>오프닝</h2><div class="quote">“데이터 보호와 계획서의 권한별 이용 기능을 함께 구현하기 위해, 각 기관의 전제를 확인하고 연결 지점과 담당 업무를 정리하고자 합니다.”</div><p class="small muted">회신이 없으면 ‘미확인’으로 기록합니다. 동의·반대로 해석하거나 특정 기관의 책임을 미리 확정하지 않습니다.</p><footer>1 / 2　기관별 질문·일정은 회의 준비 제안입니다. 원문 PDF와 1쪽 협의표를 함께 사용하세요.</footer></section>'''
    p2='''<section class="page"><span class="tag">회의 중 결정·기록용 안내</span><h1>대안 선택과 합의 후 검증</h1><h2>조건을 확인한 뒤 연결 방식 선택</h2><table><colgroup><col style="width:21%"><col style="width:24%"><col style="width:55%"></colgroup><tr><th>기술 대안</th><th>초기 비용 / 난이도</th><th>조건·한계</th></tr><tr><td>공개자료 게시</td><td>낮음 / 낮음</td><td>이 기능만으로 권한별 DT·예측 연동 등을 대체하기 어려움</td></tr><tr><td>VPN 지정 접속</td><td>낮음~중간 / 중간</td><td>기존 VPN·소수 이용자 조건 확인. 서비스별 권한 별도 필요</td></tr><tr><td>제한 API 중계</td><td>중간~높음 / 중간~높음</td><td>내부 API 접속 허용, 인증 연계·보안 운영 담당 필요</td></tr><tr><td>조회용 사본</td><td>중간~높음 / 중간~높음</td><td>외부 전송 허용. 갱신·삭제·권한 반영, 예측·제어 별도 검토</td></tr></table><p class="small muted">상대 비교이며 견적·확정 설계가 아닙니다. 이용자·데이터량·주기·기존 자원 확인 후 개발 및 운영비를 산정합니다.</p><div class="note"><b>비용 검토 시:</b> 조합의 핵심 연동·보안 R&D 자체 개발과 외주 범위(PDF 28~29쪽)를 함께 확인합니다.</div><h2>논의가 엇갈릴 때 확인할 문장</h2><ul><li>“외부→내부 접속과 내부→외부 전송 중 어느 쪽이 제한됩니까?”</li><li>“회원 인증과 최종 데이터 열람 판단이 같은 단계인지 구분하겠습니다.”</li><li>“실데이터는 어느 저장소에서 어떤 서비스로 누구에게 제공됩니까?”</li><li>“개발·서버 운영·데이터 공개 범위 확인 담당을 각각 기록하겠습니다.”</li><li>“지금 결정하기 어려운 항목은 담당기관과 확인 기한을 남기겠습니다.”</li></ul><h2>시범 연계 검증안</h2><p>비민감 샘플로 아래 항목의 담당·기한·완료 증빙을 합의합니다.</p><table><tr><th style="width:27%">검증 항목</th><th>확인할 결과</th></tr><tr><td>허용 사용자</td><td>필요한 인증 후 허용 데이터 조회 성공</td></tr><tr><td>미인증·권한 부족</td><td>보호 데이터 요청 거부</td></tr><tr><td>권한 회수</td><td>합의한 반영 시간 안에 접근 차단</td></tr><tr><td>접근 기록</td><td>사용자·시각·대상·허용/거부 결과 확인</td></tr><tr><td>예측 연계(포함 시)</td><td>입력 조건과 반환 결과의 대응 확인</td></tr><tr><td>사본 제공(채택 시)</td><td>갱신·삭제·권한 변경 반영 확인</td></tr></table><p class="small muted">위 항목은 제안이며 아직 수행하지 않았습니다. 장비 제어는 별도 대상·권한·시험 담당을 정하고, 조회 검증만으로 전체 과제 완료를 판단하지 않습니다.</p><h2>종료 전 체크</h2><p>□ 합의 문구·적용 조건　□ 개발·운영·공개 범위 확인 담당<br>□ 송수신·인증 경로　□ 미결 담당·기한　□ 검증 증빙·후속 일정</p><footer>2 / 2　근거: 연구개발계획서 PDF 21, 23, 26, 28–29, 34–35, 46, 63쪽(파일 페이지).<br>공유 자료는 계획서 내용과 중립적인 협의 질문으로 구성했습니다. 기관별 이메일 의견은 별도 확인 자료로 취급합니다.</footer></section>'''
    (ROOT/'meeting_prep.html').write_text(doc(p1+p2,'수요일 회의 준비·진행안'),encoding='utf-8')

def make_workbook():
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,headers,rows,widths):
        ws=wb.create_sheet(name);ws.append(headers)
        for row in rows:ws.append(row)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        ws.sheet_view.showGridLines=False
        for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='183249');cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF',size=11)
        ws.row_dimensions[1].height=32
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font=Font(name='맑은 고딕',size=11,color='183249');c.alignment=Alignment(vertical='top',wrap_text=True)
                c.border=Border(bottom=Side(style='thin',color='D6E1E8'))
                if c.row%2==0:c.fill=PatternFill('solid',fgColor='F0F6F9')
            ws.row_dimensions[row[0].row].height=85
        for i,w in enumerate(widths,1):ws.column_dimensions[ws.cell(1,i).column_letter].width=w
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:1';ws.print_options.horizontalCentered=True
        return ws
    sheet('사용안내',['구분','설명'],[
        ['자료 성격','계획서 기반 협의 초안. 날짜·참석·역할 배정·대안·검증은 결정 또는 수행 전입니다.'],
        ['기록 원칙','계획서 사실, 기관별 제안, 공동 합의를 구분합니다. 의견 출처·확인일과 적용 대상을 기록하세요.'],
        ['진행 순서','협의표 D1~D6 → 합의기록 → 통신경로·권한표 → 비용검토·검증체크'],
        ['기관 구분','기관명은 개발 담당이며 서버 소유·운영기관은 별도 확인합니다. 협의기관 목록은 참석 제안입니다.'],
        ['출처','RS-2026-25526544 연구개발계획서. 모든 쪽수는 PDF 파일 페이지입니다.'],
        ['배포','공유 원본은 중립적인 질문과 빈 양식입니다. 작성 후에는 회의 참여자·기관의 공유 범위에 따라 관리하세요.'],
    ],[20,110])
    sheet('협의표',['ID','항목','계획서 기재 내용','회의 질문','필요 결과','협의기관(안)'],ROWS,[9,22,45,60,39,50])
    ws=sheet('합의기록',['ID','기관 의견·출처·확인일','공동 결정문·적용 조건','상태','개발 담당','운영 담당','공개 범위 확인 담당','확인·완료 기한','검증·완료 증빙','미결사항·후속 조치'],[[r[0],'','','미확인','','','','','',''] for r in ROWS],[9,40,50,16,22,22,25,22,35,40])
    dv=DataValidation(type='list',formula1='"미확인,논의중,합의,추가확인"');ws.add_data_validation(dv);dv.add('D2:D7')
    sheet('통신경로',['ID','통신 시작 시스템','수신 시스템','사용자 요청 또는 전송 목적','데이터/API·포트','시작 방향·연결 방식','인증 확인 위치','데이터 권한 확인 위치','전송·저장·기록 담당','운영 조건·근거'],[[f'F{i:02d}','','','','','','','','',''] for i in range(1,7)],[10,25,25,35,30,33,28,28,35,40])
    sheet('권한표',['이용자 구분(검토용)','대상 데이터·등급','조회','예측','DT 열람','장비 제어','추가 인증 조건','권한 부여·회수 담당','권한 확인 시스템','합의 근거'],[[x,'','미정','미정','미정','미정','','','',''] for x in ['기업 내부 연구자','외부 과제 참여자','기타 회원','일반 방문자']],[25,35,16,16,16,16,32,30,30,35])
    ws=wb['권한표'];dv=DataValidation(type='list',formula1='"미정,허용,거부,조건부"');ws.add_data_validation(dv);dv.add('C2:F5')
    opts=[('공개자료 별도 게시','낮음','낮음~중간','낮음','단독으로 계획된 DT·예측·권한 기능 충족을 설명하기 어려움'),('VPN 지정 접속','낮음~중간','이용자 증가 시 상승','중간','기존 VPN·서비스별 권한·조합 연계 확인'),('제한 API 중계','중간~높음','중간~높음','중간~높음','내부 API 접속 조건·보안 운영 담당 확인'),('허용 데이터 사본','중간~높음','중간~높음','중간~높음','외부 전송·갱신·삭제·권한 반영 및 예측·제어 별도 확인')]
    sheet('비용검토',['대안(기술 제안)','초기비용(상대)','운영부담(상대)','난이도(상대)','조건·한계','기존 재사용 자원','추가 개발 범위·인월','일회 구축비·근거','월/연 운영비·근거','담당기관·추가 확인'],[[*o,'','','','',''] for o in opts],[24,20,25,20,55,35,35,30,30,40])
    tests=[('V1','허용 사용자','필요 인증 후 허용 샘플 데이터 조회 성공'),('V2','미인증·권한 부족','보호 데이터 요청 거부'),('V3','권한 회수','합의한 반영 시간 안에 접근 차단'),('V4','접근 기록','사용자·시각·대상·허용/거부 결과 확인'),('V5','예측 연계(포함 시)','입력 조건과 반환 결과의 대응 확인'),('V6','사본 제공(채택 시)','갱신·삭제·권한 변경 반영 확인')]
    sheet('검증체크',['ID','검증안','확인 기준(제안)','대상 시스템·샘플','담당','기한','결과','증빙·보완'],[[*t,'','','','미실시',''] for t in tests],[10,26,55,35,25,20,18,45])
    path=ROOT/'meeting_workbook.xlsx';wb.save(path)
    check=load_workbook(path);assert len(check.sheetnames)==7;assert check['합의기록']['D2'].value=='미확인'

async def render():
    async with async_playwright() as p:
        browser=await p.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox'])
        for stem,width,height,pages in [('discussion_sheet',1123,794,1),('meeting_prep',794,1122,2)]:
            page=await browser.new_page(viewport={'width':width,'height':height},device_scale_factor=2)
            await page.set_content((ROOT/f'{stem}.html').read_text(),wait_until='load');await page.evaluate('document.fonts.ready')
            await page.pdf(path=str(ROOT/f'{stem}.pdf'),width=f'{width}px',height=f'{height}px',print_background=True,margin={'top':'0','bottom':'0','left':'0','right':'0'})
            if stem=='discussion_sheet':await page.screenshot(path=str(ROOT/f'{stem}.png'),full_page=True)
            pdf=fitz.open(ROOT/f'{stem}.pdf');assert len(pdf)==pages,(stem,len(pdf))
            assert all('\ufffd' not in x.get_text() for x in pdf)
            print(stem, 'pages=',len(pdf))
            await page.close()
        await browser.close()

def finish():
    readme='''# 수요일 회의 준비자료

계획서의 기능과 기관별 구현 조건을 대조해 **적용 범위·연결 방식·담당·검증 기준**을 결정하기 위한 자료입니다. 회의 일시·방식·참석기관은 미정입니다.

- [1쪽 협의표 PDF](discussion_sheet.pdf) · [이미지](discussion_sheet.png)
- [회의 기록용 엑셀](meeting_workbook.xlsx): 협의표, 합의기록, 통신경로, 권한표, 비용검토, 검증체크
- [준비·진행안 PDF (2쪽)](meeting_prep.pdf) · [본문으로 읽기](meeting_prep.md)
- [전체 자료 ZIP](meeting_preparation.zip)

![기관별 역할·연계 협의표](discussion_sheet.png)

질문·참여기관·진행 시간·검증 기준·기술 대안은 회의 제안입니다. 기관 간 합의 또는 시험 완료를 뜻하지 않습니다. 기존 메일의 원문·연락처·미확정 기관 입장은 포함하지 않았습니다.

근거: RS-2026-25526544 연구개발계획서 PDF 파일 페이지 21, 23, 26, 28–29, 34–35, 46, 63쪽. 원문과 회의 결과를 함께 확인하세요.
'''
    (ROOT/'README.md').write_text(readme,encoding='utf-8')
    with zipfile.ZipFile(ROOT/'meeting_preparation.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(ROOT.iterdir()):
            if p.is_file() and p.suffix!='.zip':z.write(p,'meeting_preparation/'+p.name)
    with zipfile.ZipFile(ROOT/'meeting_preparation.zip') as z:assert z.testzip() is None

if __name__=='__main__':
    build_sheet();build_guide();make_workbook();asyncio.run(render());finish()
